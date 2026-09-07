"""
Document Exporter for ChurnGuard AI:
Generates two separate publication-grade documents in both PDF and DOCX formats:
1. ChurnGuard_AI_Case_Study (.pdf and .docx)
2. ChurnGuard_AI_Zero_to_Hero_Interview_Guide (.pdf and .docx)
"""

import os
import re
import base64
import subprocess
from pathlib import Path
import markdown
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_ROOT = Path(__file__).resolve().parent
DOCS_DIR = WORKSPACE_ROOT / "docs"
EXPORTS_DIR = DOCS_DIR / "exports"
EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

# ==============================================================================
# CSS Theme for PDF Generation
# ==============================================================================
PDF_CSS = """
@page {
    size: A4;
    margin: 18mm 16mm 18mm 16mm;
    @bottom-right {
        content: counter(page);
    }
}

body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #1e293b;
    line-height: 1.6;
    font-size: 11pt;
    margin: 0;
    padding: 0;
}

h1 {
    color: #0f172a;
    font-size: 22pt;
    font-weight: 800;
    border-bottom: 2.5px solid #0f172a;
    padding-bottom: 8px;
    margin-top: 0;
    margin-bottom: 16px;
    letter-spacing: -0.5px;
}

h2 {
    color: #0f172a;
    font-size: 15pt;
    font-weight: 700;
    border-bottom: 1px solid #cbd5e1;
    padding-bottom: 6px;
    margin-top: 24px;
    margin-bottom: 12px;
    page-break-after: avoid;
}

h3 {
    color: #1e293b;
    font-size: 12.5pt;
    font-weight: 700;
    margin-top: 18px;
    margin-bottom: 8px;
    page-break-after: avoid;
}

h4 {
    color: #334155;
    font-size: 11pt;
    font-weight: 600;
    margin-top: 14px;
    margin-bottom: 6px;
}

p, li {
    color: #334155;
    font-size: 10.5pt;
    margin-top: 0;
    margin-bottom: 8px;
}

ul, ol {
    margin-top: 4px;
    margin-bottom: 12px;
    padding-left: 24px;
}

li {
    margin-bottom: 4px;
}

strong {
    color: #0f172a;
    font-weight: 600;
}

code {
    font-family: Consolas, "Liberation Mono", Menlo, Courier, monospace;
    background-color: #f1f5f9;
    color: #0f172a;
    padding: 2px 5px;
    border-radius: 4px;
    font-size: 9.5pt;
    border: 1px solid #e2e8f0;
}

pre {
    background-color: #0f172a;
    color: #e2e8f0;
    padding: 12px 16px;
    border-radius: 6px;
    overflow-x: auto;
    font-size: 9pt;
    line-height: 1.45;
    margin-top: 8px;
    margin-bottom: 16px;
    page-break-inside: avoid;
}

pre code {
    background-color: transparent;
    color: #e2e8f0;
    padding: 0;
    border: none;
    font-size: 9pt;
}

table {
    width: 100%;
    border-collapse: collapse;
    margin-top: 12px;
    margin-bottom: 18px;
    font-size: 9.5pt;
    page-break-inside: avoid;
}

th, td {
    padding: 8px 10px;
    text-align: left;
    border: 1px solid #cbd5e1;
}

th {
    background-color: #0f172a;
    color: #ffffff;
    font-weight: 600;
}

tr:nth-child(even) {
    background-color: #f8fafc;
}

blockquote {
    border-left: 4px solid #3b82f6;
    background-color: #f8fafc;
    margin: 12px 0;
    padding: 10px 16px;
    color: #334155;
    font-style: normal;
    border-radius: 0 6px 6px 0;
    page-break-inside: avoid;
}

blockquote p {
    margin: 0;
}

img {
    max-width: 95%;
    height: auto;
    display: block;
    margin: 14px auto;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
    box-shadow: 0 4px 10px rgba(0, 0, 0, 0.06);
    page-break-inside: avoid;
}

.doc-header-badge {
    display: inline-block;
    background-color: #0f172a;
    color: #ffffff;
    font-size: 9pt;
    font-weight: 700;
    padding: 4px 10px;
    border-radius: 4px;
    margin-bottom: 8px;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}

.doc-meta {
    font-size: 9.5pt;
    color: #64748b;
    margin-bottom: 20px;
    border-bottom: 1px solid #e2e8f0;
    padding-bottom: 10px;
}
"""


def convert_image_tags_to_base64(html_content: str, base_dir: Path) -> str:
    """Replaces relative image src paths with embedded base64 data URIs for robust rendering."""
    def replacer(match):
        img_src = match.group(1)
        # Check potential paths
        candidate_paths = [
            base_dir / img_src,
            WORKSPACE_ROOT / img_src,
            WORKSPACE_ROOT / "docs" / img_src,
            WORKSPACE_ROOT / "docs" / "screenshots" / Path(img_src).name,
        ]
        found_path = None
        for p in candidate_paths:
            if p.exists() and p.is_file():
                found_path = p
                break

        if found_path:
            ext = found_path.suffix.lower().replace(".", "")
            mime = "image/png" if ext == "png" else "image/jpeg"
            with open(found_path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode("utf-8")
            return f'src="data:{mime};base64,{b64}"'
        return match.group(0)

    return re.sub(r'src=["\']([^"\']+)["\']', replacer, html_content)


def generate_pdf_via_chrome(html_path: Path, output_pdf_path: Path):
    """Uses Puppeteer/headless Chrome to generate pixel-perfect A4 PDF."""
    runner_script = WORKSPACE_ROOT / "frontend" / "temp_pdf_runner.cjs"
    script_content = f"""
const puppeteer = require('puppeteer-core');
const path = require('path');

(async () => {{
  const browser = await puppeteer.launch({{
    executablePath: '{CHROME_PATH.replace(chr(92), "/")}',
    headless: 'new',
    args: ['--no-sandbox', '--disable-setuid-sandbox']
  }});
  const page = await browser.newPage();
  await page.goto('file:///' + path.resolve('{str(html_path).replace(chr(92), "/")}'), {{ waitUntil: 'networkidle0' }});
  await page.pdf({{
    path: path.resolve('{str(output_pdf_path).replace(chr(92), "/")}'),
    format: 'A4',
    printBackground: true,
    margin: {{
      top: '18mm',
      bottom: '18mm',
      left: '16mm',
      right: '16mm'
    }},
    displayHeaderFooter: true,
    headerTemplate: '<div></div>',
    footerTemplate: '<div style="font-family: sans-serif; font-size: 8pt; text-align: right; width: 100%; padding-right: 16mm; color: #64748b;">Page <span class="pageNumber"></span> of <span class="totalPages"></span></div>'
  }});
  await browser.close();
  console.log('PDF generated: {output_pdf_path.name}');
}})();
"""
    with open(runner_script, "w", encoding="utf-8") as f:
        f.write(script_content)

    subprocess.run(["node", str(runner_script)], cwd=WORKSPACE_ROOT / "frontend", check=True)
    if runner_script.exists():
        runner_script.unlink()


def create_styled_html(md_path: Path, badge_text: str) -> str:
    """Renders markdown file to a complete, standalone styled HTML document."""
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    # Clean markdown alerts for standard rendering
    md_text = re.sub(r'>\s*\[!NOTE\]', '> **NOTE:**', md_text)
    md_text = re.sub(r'>\s*\[!IMPORTANT\]', '> **IMPORTANT:**', md_text)
    md_text = re.sub(r'>\s*\[!TIP\]', '> **TIP:**', md_text)
    md_text = re.sub(r'>\s*\[!WARNING\]', '> **WARNING:**', md_text)

    html_body = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "nl2br", "sane_lists"]
    )

    html_body = convert_image_tags_to_base64(html_body, md_path.parent)

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{md_path.stem}</title>
    <style>
        {PDF_CSS}
    </style>
</head>
<body>
    <div class="doc-header-badge">{badge_text}</div>
    {html_body}
</body>
</html>
"""
    return full_html


# ==============================================================================
# DOCX Generation using python-docx
# ==============================================================================
def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def style_docx_document(doc):
    """Sets standard styling for Word document."""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 65, 85)


def build_docx_from_markdown(md_path: Path, output_docx_path: Path, doc_title: str, badge: str):
    """Builds a rich Word (.docx) document from markdown content."""
    doc = docx.Document()
    style_docx_document(doc)

    # Header Badge
    badge_p = doc.add_paragraph()
    badge_run = badge_p.add_run(f"[{badge}]")
    badge_run.bold = True
    badge_run.font.size = Pt(9.5)
    badge_run.font.color.rgb = RGBColor(15, 23, 42)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    for line in lines:
        stripped = line.rstrip()

        # Handle Code blocks
        if stripped.startswith("```"):
            if in_code_block:
                # Flush code block
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(stripped)
            continue

        # Handle Markdown Tables
        if "|" in stripped and not in_code_block:
            parts = [p.strip() for p in stripped.split("|")[1:-1]]
            if len(parts) >= 2:
                # Check if separator line
                if all(re.match(r'^:?-+:?$', p) for p in parts):
                    continue
                table_rows.append(parts)
                in_table = True
                continue

        if in_table and ("|" not in stripped or stripped == ""):
            # Render accumulated table
            if table_rows:
                ncols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=ncols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(table_rows):
                    row = t.rows[r_idx]
                    for c_idx, val in enumerate(row_data):
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            cell.text = val
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            if r_idx == 0:
                                set_cell_background(cell, "0F172A")
                                for run in p.runs:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                    run.font.size = Pt(9.5)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC")
                                for run in p.runs:
                                    run.font.size = Pt(9.5)
                doc.add_paragraph()  # spacer
            table_rows = []
            in_table = False

        if not stripped:
            continue

        # Handle Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[2:].strip())
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            continue
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[3:].strip())
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            continue
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            run = h.add_run(stripped[4:].strip())
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            continue
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            run = h.add_run(stripped[5:].strip())
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            continue

        # Handle Images
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text, img_rel = img_match.groups()
            candidate_paths = [
                WORKSPACE_ROOT / img_rel,
                WORKSPACE_ROOT / "docs" / img_rel,
                WORKSPACE_ROOT / "docs" / "screenshots" / Path(img_rel).name,
            ]
            found_img = None
            for cp in candidate_paths:
                if cp.exists() and cp.is_file():
                    found_img = cp
                    break
            if found_img:
                doc.add_paragraph()  # spacer
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(str(found_img), width=Inches(6.0))
                if alt_text:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = p_cap.add_run(f"Figure: {alt_text}")
                    cap_run.font.italic = True
                    cap_run.font.size = Pt(9.5)
                    cap_run.font.color.rgb = RGBColor(100, 116, 139)
                doc.add_paragraph()  # spacer
                continue

        # Handle Blockquotes / Alerts
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            clean_bq = stripped.lstrip("> ").replace("[!NOTE]", "NOTE:").replace("[!IMPORTANT]", "IMPORTANT:")
            run = p.add_run(clean_bq)
            run.font.italic = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            continue

        # Handle Bullet points
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            text_content = stripped[2:].strip()
            # Parse bold elements inside bullets
            parts = re.split(r'(\*\*.*?\*\*)', text_content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    p.add_run(part)
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            idx, text_content = num_match.groups()
            p = doc.add_paragraph(style='List Number')
            parts = re.split(r'(\*\*.*?\*\*)', text_content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    p.add_run(part)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                p.add_run(part)

    # Save Word file
    doc.save(output_docx_path)
    print(f"DOCX generated: {output_docx_path.name}")


def main():
    print("======================================================================")
    print("STARTING DOCUMENT EXPORTS: PDF & DOCX")
    print("======================================================================")

    # 1. Export Case Study & Walkthrough
    readme_path = WORKSPACE_ROOT / "README.md"
    pdf1_path = EXPORTS_DIR / "ChurnGuard_AI_Case_Study.pdf"
    docx1_path = EXPORTS_DIR / "ChurnGuard_AI_Case_Study.docx"
    html1_path = EXPORTS_DIR / "temp_case_study.html"

    print("\n--- 1. Generating Case Study & Walkthrough ---")
    html1_content = create_styled_html(readme_path, "Portfolio Case Study & Technical Walkthrough")
    with open(html1_path, "w", encoding="utf-8") as f:
        f.write(html1_content)
    generate_pdf_via_chrome(html1_path, pdf1_path)
    build_docx_from_markdown(readme_path, docx1_path, "ChurnGuard AI Case Study", "CASE STUDY & WALKTHROUGH")
    if html1_path.exists():
        html1_path.unlink()

    # 2. Export Zero to Hero Learning Guide
    guide_path = WORKSPACE_ROOT / "docs" / "PROJECT_WALKTHROUGH.md"
    pdf2_path = EXPORTS_DIR / "ChurnGuard_AI_Zero_to_Hero_Interview_Guide.pdf"
    docx2_path = EXPORTS_DIR / "ChurnGuard_AI_Zero_to_Hero_Interview_Guide.docx"
    html2_path = EXPORTS_DIR / "temp_interview_guide.html"

    print("\n--- 2. Generating Zero-to-Hero Interview Guide ---")
    html2_content = create_styled_html(guide_path, "Zero-to-Hero Teaching & Interview Defense Guide")
    with open(html2_path, "w", encoding="utf-8") as f:
        f.write(html2_content)
    generate_pdf_via_chrome(html2_path, pdf2_path)
    build_docx_from_markdown(guide_path, docx2_path, "ChurnGuard AI Learning Guide", "INTERVIEW STUDY GUIDE")
    if html2_path.exists():
        html2_path.unlink()

    # Copy exports to root folder as well for immediate accessibility
    for src in [pdf1_path, docx1_path, pdf2_path, docx2_path]:
        dst = WORKSPACE_ROOT / src.name
        with open(src, "rb") as sf, open(dst, "wb") as df:
            df.write(sf.read())

    print("\n======================================================================")
    print("ALL 4 FILES GENERATED AND SAVED IN ROOT AND DOCS/EXPORTS:")
    print(f"1. {pdf1_path.name}  ({pdf1_path.stat().st_size / 1024:.1f} KB)")
    print(f"2. {docx1_path.name} ({docx1_path.stat().st_size / 1024:.1f} KB)")
    print(f"3. {pdf2_path.name}  ({pdf2_path.stat().st_size / 1024:.1f} KB)")
    print(f"4. {docx2_path.name} ({docx2_path.stat().st_size / 1024:.1f} KB)")
    print("======================================================================")


if __name__ == "__main__":
    main()
